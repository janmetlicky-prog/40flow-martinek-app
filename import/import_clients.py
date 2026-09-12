#!/usr/bin/env python3
"""
Import listu „Klienti" z Excelu do data/clients.json.

Opakovaně spustitelný — přepíše clients.json novým stavem z Excelu.
Existující JSON předtím zálohuje do output/clients_backup_<timestamp>.json.

Použití:
    python3 import/import_clients.py /cesta/k/Obchody.xlsx
"""
from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_FILE = PROJECT_ROOT / "data" / "clients.json"
BACKUP_DIR = PROJECT_ROOT / "output"

SHEET = "Klienti"
HEADER_ROW = 2  # řádek 1 = nápověda, řádek 2 = názvy sloupců, data od řádku 3

COLUMNS = [
    "stav", "jmeno", "prijmeni", "firma", "obchodnik", "dohoda", "datum",
    "lead_agent", "oblacek", "bilance", "zivot", "investice", "nezivot",
    "poznamka_nzp", "uver", "uver_bydleni", "sdileni", "poznamky_lenka",
]

PRODUCT_FIELDS = ["zivot", "investice", "nezivot", "uver", "uver_bydleni"]
PRODUCT_STATES = [
    "Budeme řešit", "Audit", "Nebudeme řešit", "Pozdějí",
    "Rozpracováno", "Předáno", "Hotovo",
]


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def import_clients(excel_path: Path) -> dict:
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"Chyba: list '{SHEET}' v souboru není. Nalezeno: {wb.sheetnames}")
    ws = wb[SHEET]

    clients = []
    skipped = 0
    for idx, row in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True), start=HEADER_ROW + 1):
        values = [_cell(v) for v in (row[:18] + (None,) * max(0, 18 - len(row)))]
        # klient bez jména, příjmení i firmy = prázdný řádek
        if not (values[1] or values[2] or values[3]):
            skipped += 1
            continue
        client = dict(zip(COLUMNS, values))
        client["id"] = f"c{idx:04d}"
        clients.append(client)

    unknown_states = sorted({
        c[f] for c in clients for f in PRODUCT_FIELDS
        if c[f] and c[f] not in PRODUCT_STATES
    })

    return {
        "meta": {
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": excel_path.name,
            "client_count": len(clients),
            "product_states": PRODUCT_STATES,
            "unknown_states_seen": unknown_states,
        },
        "clients": clients,
    }


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("Použití: python3 import/import_clients.py <soubor.xlsx>")
    excel_path = Path(sys.argv[1])
    if not excel_path.exists():
        sys.exit(f"Chyba: soubor nenalezen: {excel_path}")

    if DATA_FILE.exists():
        BACKUP_DIR.mkdir(exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = BACKUP_DIR / f"clients_backup_{stamp}.json"
        shutil.copy2(DATA_FILE, backup)
        print(f"Záloha: {backup.relative_to(PROJECT_ROOT)}")

    data = import_clients(excel_path)
    DATA_FILE.parent.mkdir(exist_ok=True)
    DATA_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Importováno {data['meta']['client_count']} klientů → {DATA_FILE.relative_to(PROJECT_ROOT)}")
    if data["meta"]["unknown_states_seen"]:
        print(f"Pozn.: hodnoty mimo číselník produktových stavů: {data['meta']['unknown_states_seen']}")


if __name__ == "__main__":
    main()
